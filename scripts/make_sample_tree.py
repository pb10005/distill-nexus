#!/usr/bin/env python3
# @covers AC-070, AC-071, AC-072
"""Generate ``tests/fixtures/sample_tree/`` (about 40 files, identical on every OS).

Windows reserved names and >260-char paths are NOT committed (they would break a
Windows checkout); tests create those in a temporary directory instead.
@assumption AS-021

usage: python scripts/make_sample_tree.py [target_dir]
"""

from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DEFAULT = ROOT / "tests" / "fixtures" / "sample_tree"

MD = {
    "inbox/terms-of-service.md": "# Terms of Service\n\nIssued 2024-01-15.\n\n## Scope\n\nThese terms govern use of the Nexus API gateway.\n\n## Liability\n\nThe provider is not liable for indirect damages.\n",
    "specs/api-spec.md": "# Nexus API Specification\n\n## Overview\n\nThe API gateway routes client requests to backend services.\n\n## Rate limit\n\nThe rate limit is 100 requests per second per client.\n\n## Authentication\n\nClients authenticate with OAuth 2.0 bearer tokens.\n",
    "specs/adr-001-event-bus.md": "# ADR-001: Adopt an event bus\n\nDate: 2024-02-01\n\n## Decision\n\nServices communicate through the event bus instead of direct calls.\n\n## Consequences\n\nThe order service publishes OrderCreated events.\n",
    "specs/requirements.md": "# Requirements\n\n## Functional\n\nThe billing service issues invoices monthly.\n\n## Non-functional\n\nThe API gateway must answer within 200 ms at p95.\n",
    "meetings/2024-03-12-kickoff.md": "# Kickoff meeting 2024-03-12\n\nAttendees: product, engineering.\n\n## Decisions\n\nThe rate limit is 200 requests per second per client.\n\n## Actions\n\nDraft the SLA.\n",
    "meetings/retro.rst": "Sprint retro\n============\n\nWhat went well: the event bus migration.\n\nWhat to improve: on-call handover.\n",
    "knowledge/glossary.md": "# Glossary\n\n## API Gateway\n\nThe entry point that routes requests to services.\n\n## Event bus\n\nA publish/subscribe channel between services.\n\n## SLA\n\nService level agreement.\n",
    "knowledge/onboarding.md": "# Onboarding\n\n## Accounts\n\nNew engineers receive an SSO account on day one.\n\n## Tools\n\nWe use the Nexus API gateway and the event bus.\n",
    "knowledge/pricing.md": "# Pricing\n\n## Plans\n\nThe Standard plan costs 50 USD per month.\n\n## Overage\n\nRequests above the quota are billed per 1000 calls.\n",
    "knowledge/roadmap.md": "# Roadmap 2025\n\n## Q1\n\nLaunch the partner portal.\n\n## Q2\n\nAdd GraphQL to the API gateway.\n",
    "knowledge/security-policy.md": "# Security policy\n\n## Secrets\n\nSecrets are stored in the vault service.\n\n## Access\n\nProduction access requires two-person approval.\n",
    "knowledge/faq.md": "# FAQ\n\n## What is the rate limit?\n\nSee the API specification.\n\n## Who owns the gateway?\n\nThe platform team owns the API gateway.\n",
    "incidents/incident-2024-05.md": "# Incident 2024-05-02\n\n## Summary\n\nThe event bus lost messages for 12 minutes.\n\n## Root cause\n\nA broker disk filled up.\n",
    "a/b/c/readme.txt": "Deeply nested readme for the sample tree.\n",
    "data/customers.csv": "id,name,plan\n1,Acme,Standard\n2,Globex,Enterprise\n",
    "data/config.yaml": "gateway:\n  rate_limit: 100\n  timeout_ms: 200\n",
    "data/settings.json": '{"feature_flags": {"graphql": false}}\n',
    "data/schema.xml": '<schema><table name="orders"/></schema>\n',
    "data/notes.toml": '[owner]\nteam = "platform"\n',
    "logs/app.log": "2024-05-02T10:00:00 ERROR broker disk full\n2024-05-02T10:12:00 INFO recovered\n",
    "code/gateway.py": '"""Nexus API gateway entry point."""\n\n\ndef route(request):\n    """Route a request to the backend service."""\n    return request\n\n\nclass RateLimiter:\n    """Token bucket limiter: 100 requests per second."""\n',
    "code/client.ts": "// Nexus API client\nexport function call(path: string) {\n  return fetch(path);\n}\n",
    "code/main.go": "// Package main starts the gateway.\npackage main\n\nfunc main() {}\n",
    "code/lib.rs": "/// Billing helpers\npub fn invoice() {}\n",
    "code/App.java": "// Order service\npublic class App {\n  public static void main(String[] a) {}\n}\n",
    "code/Program.cs": "// Event bus consumer\nclass Program { static void Main() {} }\n",
    "code/util.js": "// small helpers\nfunction add(a, b) { return a + b; }\n",
}

HTML = """<html><head><title>Architecture</title><style>body{}</style><script>var x=1;</script></head>
<body><nav>menu</nav><h1>Architecture</h1><p>The API gateway sits in front of the order and billing services.</p>
<table><tr><th>Service</th><th>Owner</th></tr><tr><td>order</td><td>commerce</td></tr></table>
<footer>footer</footer></body></html>
"""


def _docx(path: Path, title: str, paras: list[str], table: list[list[str]] | None = None) -> None:
    import docx

    d = docx.Document()
    d.add_heading(title, level=1)
    for p in paras:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, val in enumerate(row):
                t.cell(r, c).text = val
    d.core_properties.created = datetime(2024, 1, 1)
    d.save(str(path))


def _pdf(path: Path, pages: list[str | None]) -> None:
    import pymupdf

    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        if text:
            page.insert_textbox(pymupdf.Rect(50, 50, 550, 800), text, fontsize=11)
        else:
            page.draw_rect(pymupdf.Rect(100, 100, 300, 200), color=(0, 0, 0), fill=(0.8, 0.8, 0.8))
    doc.save(str(path))
    doc.close()


def _png(path: Path) -> None:
    import pymupdf

    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 64, 32), 0)
    pix.clear_with(200)
    pix.save(str(path))


def _xlsx(path: Path, rows: int) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "sales"
    ws.append(["month", "region", "amount"])
    for i in range(1, rows):
        ws.append([f"2024-{(i % 12) + 1:02d}", ["east", "west"][i % 2], i * 10])
    wb.save(str(path))


def _pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for i, (title, body, notes) in enumerate(
        [
            ("Weekly sync", "Event bus migration done", "Mention the incident"),
            ("Next steps", "Partner portal", "Owner: platform"),
        ]
    ):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = title
        s.placeholders[1].text = body
        s.notes_slide.notes_text_frame.text = notes
        if i == 0:
            s.shapes.add_textbox(Inches(1), Inches(5), Inches(4), Inches(1)).text_frame.text = "Q2 goals"
    prs.save(str(path))


def build(target: Path) -> list[Path]:
    target.mkdir(parents=True, exist_ok=True)
    made: list[Path] = []

    def put(rel: str, data: bytes | str) -> Path:
        p = target / rel
        p.parent.mkdir(parents=True, exist_ok=True)
        if isinstance(data, str):
            p.write_bytes(data.encode("utf-8"))
        else:
            p.write_bytes(data)
        made.append(p)
        return p

    for rel, text in MD.items():
        put(rel, text)
    put("specs/architecture.html", HTML)
    put(
        "meetings/minutes_cp932.txt",
        "議事録 2024-04-01\r\n決定事項: レート制限は毎秒100リクエストとする。\r\n".encode("cp932"),
    )
    put("knowledge/用語集.md", "# 用語集\n\n## イベントバス\n\nサービス間の非同期メッセージ基盤。\n")

    def gen(rel: str, fn, *args) -> None:  # type: ignore[no-untyped-def]
        p = target / rel
        p.parent.mkdir(parents=True, exist_ok=True)
        fn(p, *args)
        made.append(p)

    gen(
        "inbox/NDA_vendor-a.docx",
        _docx,
        "Mutual NDA - Vendor A",
        ["Effective 2024-03-12.", "Confidential information must not be disclosed."],
        [["Party", "Role"], ["Nexus Inc.", "Discloser"], ["Vendor A", "Recipient"]],
    )
    gen(
        "inbox/outsourcing-agreement.pdf",
        _pdf,
        [
            "Outsourcing Agreement\n\nDated 2023-11-01. The contractor delivers the billing service. " * 3,
            "Payment is due within 30 days of invoice. " * 5,
        ],
    )
    gen(
        "inbox/scanned-contract.pdf",
        _pdf,
        ["Signed contract cover page. The agreement term is 12 months and renews automatically. " * 2, None],
    )
    gen("specs/er-diagram.png", _png)
    gen("data/sales.xlsx", _xlsx, 250)
    gen("meetings/weekly.pptx", _pptx)

    # duplicates (same bytes, different names)
    put("dup/api-spec-copy.md", MD["specs/api-spec.md"])
    put("dup/NDA copy.docx", (target / "inbox/NDA_vendor-a.docx").read_bytes())
    # non-knowledge / media / type spoofing
    put("misc/video.mp4", b"\x00\x00\x00\x18ftypmp42" + b"\x00" * 64)
    put("misc/tool.exe", b"MZ" + b"\x00" * 126)
    put("misc/looks-like-text.txt", (target / "specs/er-diagram.png").read_bytes())
    # excluded from scan
    put(".hidden-notes.md", "# hidden\n")
    put(".dnignore", "*.bak\n")
    put("old.bak", "backup\n")
    return made


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT
    files = build(out)
    print(f"wrote {len(files)} files to {out}")
