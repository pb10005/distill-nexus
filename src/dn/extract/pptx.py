# @covers AC-021
"""Presentations: text and speaker notes per slide number."""

from __future__ import annotations

from pathlib import Path

from dn.extract.base import Extracted, gfm_table


def extract(path: Path) -> Extracted:
    from pptx import Presentation

    prs = Presentation(str(path))
    out: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        out.append(f"## Slide {i}")
        for shape in slide.shapes:
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [[c.text for c in r.cells] for r in shape.table.rows]
                out.append(gfm_table(rows))
            elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    out.append(text)
        if slide.has_notes_slide:
            notes = (
                slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            )
            if notes:
                out.append(f"### Notes (slide {i})\n\n{notes}")
    return Extracted(text="\n\n".join(out) + "\n", type="pptx", pages=len(prs.slides))
